"""
סקריפט ליצירת מסמך אפיון תוכנה (SRS) כקובץ Word (.docx) עבור מערכת FitStudio.
"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    if bold:
        run.bold = True

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(text)

def generate_word_doc():
    doc = Document()

    # Title
    title = doc.add_heading('מסמך אפיון תוכנה (SRS) – מערכת FitStudio', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_paragraph(doc, 'גרסת מסמך: 1.0')
    add_paragraph(doc, 'תאריך כתיבה: אוגוסט 2026')
    add_paragraph(doc, 'תיאור המערכת: מערכת פנימית לניהול סטודיו אימונים (קרוספיט/פונקציונלי), ניהול הרשמות לאימוני WOD, ניהול מנויים, וקליטה/המרת לידים.')
    doc.add_page_break()

    # Section 1
    add_heading(doc, '1. מבוא ומטרות המערכת', 1)
    add_paragraph(doc, 'מטרת המערכת היא להחליף ניהול ידני של יומן הסטודיו במערכת ממוחשבת, אמינה ומהירה. המערכת תאפשר לבעל הסטודיו (האדמין) לנהל את מערכת השעות, לעקוב אחר תפוסת אימונים, להוסיף מנויים חדשים ולנהל מתעניינים (לידים) עד להפיכתם למנויים משלמים.')

    # Section 2
    add_heading(doc, '2. חוקים עסקיים ולוגיקת ליבה (Business Logic)', 1)
    add_heading(doc, '2.1. אימוני WOD', 2)
    add_bullet(doc, 'אורך אימון: כל אימון נמשך בדיוק 60 דקות ומתחיל בשעה עגולה.')
    add_bullet(doc, 'מגבלת משתתפים (Capacity): מקסימום 18 מתאמנים רשומים לאימון.')
    add_bullet(doc, 'כפילויות: מתאמן לא יכול להירשם לאותו אימון פעמיים.')

    add_heading(doc, '2.2. שעות פעילות הסטודיו', 2)
    add_bullet(doc, "ימים א'-ה': בוקר (06:00, 07:00, 08:00), ערב (16:00, 17:00, 18:00, 19:00, 20:00).")
    add_bullet(doc, 'סופ"ש (ו\'-ש\'): בוקר (08:00, 09:00, 10:00, 11:00).')

    add_heading(doc, '2.3. ניהול לידים והמרה', 2)
    add_bullet(doc, 'ליד חדש נכנס בסטטוס NEW.')
    add_bullet(doc, 'המרה למנוי: פעולה אטומית (Transaction) שמעדכנת סטטוס ליד ל-CONVERTED ויוצרת רשומת מנוי.')

    # Section 3
    add_heading(doc, '3. ארכיטקטורה ומודל נתונים', 1)
    add_paragraph(doc, 'המערכת בנויה בארכיטקטורת Separation of Concerns, מבוססת Python ו-SQLite. ממשק המשתמש נכתב ב-Streamlit.')
    add_heading(doc, 'טבלאות מרכזיות:', 2)
    add_bullet(doc, 'members (מנויים): מנהלת מתאמנים פעילים (שם, טלפון, אימייל).')
    add_bullet(doc, 'leads (מתעניינים): מנהלת פניות וסטטוס המרה.')
    add_bullet(doc, 'wod_classes (יומן אימונים): משבצות זמן מוגדרות מראש, מניעת כפילות חלונות זמן.')
    add_bullet(doc, 'wod_registrations (הרשמות): קישור בין מתאמן לאימון, נאכף על ידי מגבלת Capacity (עד 18).')

    # Section 4 - QA
    doc.add_page_break()
    add_heading(doc, '4. תרחישי בדיקה (QA / Test Cases)', 1)
    
    # Test Case 1
    add_heading(doc, 'TC-01: אכיפת מגבלת מקסימום משתתפים (18 איש)', 2)
    add_bullet(doc, 'פעולה: ניסיון לרשום את המנוי ה-19 לאימון ספציפי.')
    add_bullet(doc, 'תוצאה צפויה: המערכת תדחה את ההרשמה ותציג שגיאה: "האימון מלא! (18/18 משתתפים)".')

    # Test Case 2
    add_heading(doc, 'TC-02: מניעת הרשמה כפולה', 2)
    add_bullet(doc, 'פעולה: ניסיון לרשום מנוי לאימון שאליו הוא כבר רשום פעיל.')
    add_bullet(doc, 'תוצאה צפויה: המערכת תזהה כפילות (UNIQUE constraint) ותציג שגיאה: "המנוי כבר רשום לאימון זה".')

    # Test Case 3
    add_heading(doc, 'TC-03: אימות שעות פעילות', 2)
    add_bullet(doc, 'פעולה: ניסיון לפתוח משבצת אימון ביום שני בשעה 11:00 (מחוץ לחלון בוקר/ערב).')
    add_bullet(doc, 'תוצאה צפויה: המערכת תחסום את הפעולה עקב כשל בוולידציה ותתריע על חריגה משעות הסטודיו.')

    # Test Case 4
    add_heading(doc, 'TC-04: המרת ליד (Transaction Atomicity)', 2)
    add_bullet(doc, 'פעולה: ביצוע המרה לליד קיים.')
    add_bullet(doc, 'תוצאה צפויה: הליד משנה סטטוס ל-CONVERTED, רשומה חדשה וזהה נוצרת בטבלת members. ניסיון המרה חוזר של אותו ליד יידחה.')

    # Save
    doc.save('FitStudio_System_Spec.docx')
    print("הקובץ 'FitStudio_System_Spec.docx' נוצר בהצלחה בתיקיית הפרויקט!")

if __name__ == "__main__":
    generate_word_doc()