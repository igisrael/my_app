import sqlite3
import pandas as pd
from docx import Document
from rag.knowledge_base import KNOWLEDGE_BASE

def export_db_to_excel(db_path, excel_path):
    print("Exporting database to Excel...")
    try:
        conn = sqlite3.connect(db_path)
        
        # Get all tables
        tables_query = "SELECT name FROM sqlite_master WHERE type='table';"
        tables = pd.read_sql_query(tables_query, conn)
        
        with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
            for table_name in tables['name']:
                # Skip sqlite internal tables
                if table_name.startswith('sqlite_'):
                    continue
                df = pd.read_sql_query(f"SELECT * from {table_name}", conn)
                df.to_excel(writer, sheet_name=table_name, index=False)
                print(f" - Exported table: {table_name}")
                
        conn.close()
        print(f"Successfully created {excel_path}")
    except Exception as e:
        print(f"Error exporting database: {e}")

def export_kb_to_word(kb_list, word_path):
    print("Exporting knowledge base to Word...")
    try:
        doc = Document()
        doc.add_heading('FitStudio - מאגר ידע לשאלות (Knowledge Base)', 0)
        
        for item in kb_list:
            q = item.get("question", "")
            a = item.get("answer", "")
            
            doc.add_heading('שאלה:', level=2)
            doc.add_paragraph(q)
            
            doc.add_heading('תשובה:', level=3)
            doc.add_paragraph(a)
            
            doc.add_paragraph("_" * 50) # separator
                
        doc.save(word_path)
        print(f"Successfully created {word_path}")
    except Exception as e:
        print(f"Error exporting knowledge base: {e}")

if __name__ == "__main__":
    db_file = "studio.db"
    excel_file = "FitStudio_Database.xlsx"
    word_file = "FitStudio_KnowledgeBase.docx"
    
    export_db_to_excel(db_file, excel_file)
    export_kb_to_word(KNOWLEDGE_BASE, word_file)
